"""Complete Person 1 data pipeline. Raw Excel files are never modified."""
from __future__ import annotations
import argparse
import re
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CASE_IDS = set(range(1, 101)); GRID_ROWS = 90_365; GRID_COLS = 29; SENTINEL = -999.25
UNCERTAINTIES = {"Fault Transmissibility":"fault_transmissibility","Porosity Multiplier":"porosity_multiplier","Permeability Multiplier":"permeability_multiplier","Aquifer Pore Volume":"aquifer_pore_volume"}
FIELD = {"Identifier":"identifier","Date":"date","Oil production rate (STB/d)":"oil_rate_stb_d","Gas production rate (MSCF/d)":"gas_rate_mscf_d","Water production rate (STB/d)":"water_rate_stb_d","Water injection rate (STB/d)":"water_injection_rate_stb_d"}
CURVES = {"Gas production cumulative":("gas","cumulative"),"Gas production rate":("gas","rate"),"Oil production cumulative":("oil","cumulative"),"Oil production rate":("oil","rate"),"Water production cumulative":("water","cumulative"),"Water production rate":("water","rate")}
FEATURES = ["Pore Volume (PORV)","Porosity (PORO)","Permeability I (PERMX)","Permeability J (PERMY)","Permeability K (PERMZ)","Net-to-Gross Ratio (NTG)","Transmissibility I (TRANX)","Transmissibility J (TRANY)","Transmissibility K (TRANZ)","Pressure (PRESSURE)","Water Saturation (SWAT)","Oil Saturation (SOIL)","Gas Saturation (SGAS)","Solution Gas-Oil Ratio (RS)","Vapor Oil-Gas Ratio (RV)"]

def output_path(processed: Path, name: str) -> Path:
    """Return the official Person 1 path for a generated CSV output."""
    return processed / f"person1_{name}"

def normalise_case_id(value: object) -> int:
    """Convert a source label such as ``Case 17`` into integer case ID 17."""
    match = re.fullmatch(r"\s*Case\s+(\d+)\s*", str(value))
    if not match: raise ValueError(f"Unexpected case label: {value!r}")
    return int(match.group(1))

def parse_case_header(value: object) -> int | None:
    """Return an ID from curve header ``Case_17``; ignore non-case headers."""
    match = re.fullmatch(r"Case_(\d+)", str(value)); return int(match.group(1)) if match else None

def split_for_case(case_id: int) -> str:
    """Assign the fixed competition split: train, validation, or test."""
    if 1 <= case_id <= 70: return "train"
    if 71 <= case_id <= 85: return "validation"
    if 86 <= case_id <= 100: return "test"
    raise ValueError(f"Unexpected case_id: {case_id}")

def grid_case_id(name: str) -> int:
    """Validate and extract the case ID from a grid worksheet name."""
    match = re.fullmatch(r"Case\s+(\d+)", name.strip())
    if not match: raise ValueError(f"Unexpected grid sheet name: {name!r}")
    return int(match.group(1))

def require_columns(frame: pd.DataFrame, required: set[str], source: str) -> None:
    """Fail early when an input table does not contain its required schema."""
    missing = required - set(frame.columns)
    if missing: raise ValueError(f"{source} is missing columns: {sorted(missing)}")

def core_audit(unc: pd.DataFrame, hist: pd.DataFrame) -> pd.DataFrame:
    """Create transparent QA checks for uncertainty inputs and field history."""
    rates = [x for x in hist if x.endswith("_d")]
    rows = [("case_uncertainties","row_count",len(unc),100,"error"),("case_uncertainties","unique_case_ids",unc.case_id.nunique(),100,"error"),("case_uncertainties","case_ids_1_to_100",set(unc.case_id)==CASE_IDS,True,"error"),("case_uncertainties","missing_values",int(unc.isna().sum().sum()),0,"error"),("field_history","row_count",len(hist),41,"warning"),("field_history","duplicate_dates",int(hist.date.duplicated().sum()),0,"error"),("field_history","missing_values",int(hist.isna().sum().sum()),0,"error"),("field_history","negative_rates",int((hist[rates]<0).sum().sum()),0,"warning")]
    audit = pd.DataFrame(rows,columns=["dataset","check","observed","expected","severity"]); audit["status"] = np.where(audit.observed.astype(str).eq(audit.expected.astype(str)),"pass","flag"); return audit

def run_core(input_dir: Path, processed: Path) -> pd.DataFrame:
    """Clean Excel 01 uncertainties and Excel 02 field history; save CSV/audit outputs."""
    intro, production = input_dir/"01 Introduction.xlsx", input_dir/"02 Field Production.xlsx"
    raw_unc = pd.read_excel(intro,sheet_name="Uncertainty Values",engine="openpyxl"); raw_hist = pd.read_excel(production,sheet_name="Field Production",engine="openpyxl")
    require_columns(raw_unc,{"Item",*UNCERTAINTIES},"Uncertainty Values"); require_columns(raw_hist,set(FIELD),"Field Production")
    unc = raw_unc[["Item",*UNCERTAINTIES]].rename(columns={"Item":"case_label",**UNCERTAINTIES}).copy(); unc.insert(0,"case_id",unc.pop("case_label").map(normalise_case_id)); unc = unc.sort_values("case_id")
    hist = raw_hist[list(FIELD)].rename(columns=FIELD).copy(); hist.date = pd.to_datetime(hist.date,errors="coerce")
    for col in set(FIELD.values())-{"identifier","date"}: hist[col] = pd.to_numeric(hist[col],errors="coerce")
    hist = hist.sort_values("date"); audit = core_audit(unc,hist); processed.mkdir(parents=True,exist_ok=True)
    unc.to_csv(output_path(processed,"case_uncertainties.csv"),index=False); hist.to_csv(output_path(processed,"field_history.csv"),index=False,date_format="%Y-%m-%d"); audit.to_csv(output_path(processed,"core_cleaning_audit.csv"),index=False)
    pd.DataFrame([{"source_file":intro.name,"sheet":"Uncertainty Values","rows":len(unc),"columns":len(unc.columns)},{"source_file":production.name,"sheet":"Field Production","rows":len(hist),"columns":len(hist.columns)}]).to_csv(output_path(processed,"core_source_inventory.csv"),index=False)
    return audit

def extract_curves(source: Path) -> tuple[pd.DataFrame,pd.DataFrame]:
    """Read six production sheets into tidy case curves and separate RAMP history."""
    cases=[]; ramp=[]; book=load_workbook(source,read_only=True,data_only=True)
    for sheet_name,(fluid,metric) in CURVES.items():
        sheet=book[sheet_name]; headers=[cell.value for cell in sheet[2]]; columns={i:parse_case_header(x) for i,x in enumerate(headers)}; ramp_cols={i for i,x in enumerate(headers) if x=="RAMP_History"}
        if {x for x in columns.values() if x is not None} != CASE_IDS: raise ValueError(f"{sheet_name} is missing cases.")
        for row in sheet.iter_rows(min_row=4,values_only=True):
            if not row or row[0] is None: continue
            for i,value in enumerate(row[1:],1):
                if value is None: continue
                if columns.get(i) is not None: cases.append({"date":pd.Timestamp(row[0]),"case_id":columns[i],"fluid":fluid,"metric":metric,"value":float(value)})
                elif i in ramp_cols: ramp.append({"date":pd.Timestamp(row[0]),"fluid":fluid,"metric":metric,"value":float(value)})
    book.close(); return pd.DataFrame(cases),pd.DataFrame(ramp)

def run_curves(input_dir: Path, processed: Path) -> pd.DataFrame:
    """Validate, quarterly-align, and save simulated production curves without extrapolation."""
    curves,ramp=extract_curves(input_dir/"02 Field Production.xlsx"); keys=["case_id","date","fluid","metric"]
    if curves.groupby(keys).value.nunique().gt(1).any(): raise ValueError("Conflicting duplicate curve records.")
    clean=curves.drop_duplicates(keys).sort_values(keys); cumulative=clean[clean.metric.eq("cumulative")].copy(); decreases=int((cumulative.groupby(["case_id","fluid"]).value.diff() < -1e-9).sum())
    checks=[("case_ids",clean.case_id.nunique(),100,"error"),("fluids",clean.fluid.nunique(),3,"error"),("metrics",clean.metric.nunique(),2,"error"),("missing_values",int(clean.isna().sum().sum()),0,"error"),("negative_values",int((clean.value<0).sum()),0,"warning"),("duplicate_records",int(curves.duplicated(keys,keep=False).sum()),0,"warning"),("decreasing_cumulative_values",decreases,0,"warning")]
    audit=pd.DataFrame(checks,columns=["check","observed","expected","severity"]); audit["status"]=np.where(audit.observed.astype(str).eq(audit.expected.astype(str)),"pass","flag")
    hist=pd.read_csv(output_path(processed,"field_history.csv"),parse_dates=["date"]); template=pd.read_excel(input_dir/"09 Template Deliverable.xlsx",sheet_name="Case 1",usecols="A",engine="openpyxl").iloc[:,0]
    dates=pd.DatetimeIndex(pd.concat([hist.date,pd.to_datetime(template,errors="coerce").dropna()]).drop_duplicates().sort_values()); target=dates.astype("int64").to_numpy(); frames=[]
    for (case,fluid,metric),group in clean.groupby(["case_id","fluid","metric"]):
        group=group.sort_values("date"); source=group.date.astype("int64").to_numpy(); inside=(target>=source.min())&(target<=source.max()); values=np.full(len(dates),np.nan); values[inside]=np.interp(target[inside],source,group.value.to_numpy(float)); status=np.where(~inside,"outside_source_range",np.where(np.isin(target,source),"exact","linear_interpolation")); frames.append(pd.DataFrame({"case_id":case,"date":dates,"fluid":fluid,"metric":metric,"value":values,"source_status":status}))
    long=pd.concat(frames,ignore_index=True); wide=long.pivot(index=["case_id","date","fluid"],columns="metric",values="value").reset_index().sort_values(["case_id","date","fluid"]); coverage=long.groupby("source_status").size().rename("records").reset_index()
    clean.to_csv(output_path(processed,"case_production_raw_long.csv"),index=False); ramp.to_csv(output_path(processed,"ramp_history_raw_long.csv"),index=False); long.to_csv(output_path(processed,"case_production_quarterly_long.csv"),index=False); wide.to_csv(output_path(processed,"case_production_quarterly.csv"),index=False); coverage.to_csv(output_path(processed,"case_production_quarterly_coverage.csv"),index=False); audit.to_csv(output_path(processed,"production_curve_cleaning_audit.csv"),index=False); return audit

def sources(input_dir: Path) -> list[Path]:
    """List the six reservoir-grid workbooks and exclude Excel temporary lock files."""
    return sorted(p for p in input_dir.glob("*.xlsx") if p.name.startswith(("03 ","04 ","05 ","06 ","07 ","08 ")) and not p.name.startswith("~$"))

def manifest(input_dir: Path, processed: Path) -> tuple[pd.DataFrame,pd.DataFrame]:
    """Inventory all grid sheets and verify their case IDs, row count, and schema width."""
    rows=[]
    for path in sources(input_dir):
        book=load_workbook(path,read_only=True,data_only=True)
        for name in book.sheetnames:
            sheet=book[name]; cid=grid_case_id(name); rows.append({"source_file":path.name,"sheet_name":name,"case_id":cid,"split":split_for_case(cid),"data_rows":sheet.max_row-1,"columns":sheet.max_column,"schema_valid":sheet.max_row-1==GRID_ROWS and sheet.max_column==GRID_COLS})
        book.close()
    table=pd.DataFrame(rows).sort_values("case_id"); checks=[("case_sheets",len(table),100),("case_ids_1_to_100",set(table.case_id)==CASE_IDS,True),("data_rows_per_case",table.data_rows.nunique(),1),("expected_data_rows",int(table.data_rows.iloc[0]),GRID_ROWS),("expected_columns",int(table["columns"].iloc[0]),GRID_COLS),("valid_schemas",int(table.schema_valid.sum()),100)]; audit=pd.DataFrame(checks,columns=["check","observed","expected"]); audit["status"]=np.where(audit.observed.astype(str).eq(audit.expected.astype(str)),"pass","flag"); table.to_csv(output_path(processed,"grid_manifest.csv"),index=False); audit.to_csv(output_path(processed,"grid_schema_audit.csv"),index=False); return table,audit

def grid_summary(path: Path,name: str) -> tuple[dict,dict]:
    """Stream one grid sheet into sentinel-safe case features plus its QA record."""
    book=load_workbook(path,read_only=True,data_only=True); sheet=book[name]; rows=sheet.iter_rows(values_only=True); pos={re.sub(r"\s+"," ",str(x or "")).strip():i for i,x in enumerate(next(rows))}; missing=[x for x in FEATURES if x not in pos]
    if missing: book.close(); raise ValueError(f"{path.name}/{name} missing {missing}")
    ix=[pos[x] for x in FEATURES]; sums=np.zeros(15); squares=np.zeros(15); count=np.zeros(15,int); low=np.full(15,np.inf); high=np.full(15,-np.inf); n=missing_count=sentinel_count=inactive=0
    for row in rows:
        values=np.array([row[i] for i in ix],float); finite=np.isfinite(values); sent=values==SENTINEL; valid=finite&~sent; n+=1; missing_count+=int((~finite).sum()); sentinel_count+=int(sent.sum()); inactive+=int(sent[1:].all()); sums+=np.where(valid,values,0); squares+=np.where(valid,values*values,0); count+=valid; low=np.minimum(low,np.where(valid,values,np.inf)); high=np.maximum(high,np.where(valid,values,-np.inf))
    book.close()
    if (count==0).any(): raise ValueError(f"{path.name}/{name} contains an all-sentinel feature.")
    cid=grid_case_id(name); mean=sums/count; std=np.sqrt(np.maximum(squares/count-mean*mean,0)); feat={"case_id":cid,"split":split_for_case(cid),"source_file":path.name,"source_sheet":name,"active_cell_count":n-inactive,"inactive_cell_count":inactive,"active_cell_fraction":(n-inactive)/n}
    for i,col in enumerate(FEATURES):
        short=re.search(r"\(([^)]+)\)",col).group(1).lower(); feat.update({f"{short}_mean":mean[i],f"{short}_std":std[i],f"{short}_min":low[i],f"{short}_max":high[i]})
    audit={"case_id":cid,"split":split_for_case(cid),"source_file":path.name,"source_sheet":name,"grid_rows":n,"feature_columns":15,"source_missing_feature_values":missing_count,"sentinel_feature_values":sentinel_count,"inactive_rows":inactive,"active_rows":n-inactive,"active_cell_fraction":(n-inactive)/n,"status":"pass" if n==GRID_ROWS and missing_count==0 else "review"}; return feat,audit

def extract_features(input_dir: Path,processed: Path,resume=False) -> tuple[pd.DataFrame,pd.DataFrame]:
    """Create one feature row per case, optionally resuming verified prior extraction."""
    fp,ap=output_path(processed,"grid_case_features.csv"),output_path(processed,"grid_feature_audit.csv"); old=pd.read_csv(fp) if resume and fp.exists() else pd.DataFrame(); features=old.to_dict("records") if not old.empty else []; audits=pd.read_csv(ap).to_dict("records") if resume and ap.exists() else []; done=set(old.case_id) if not old.empty else set()
    for path in sources(input_dir):
        book=load_workbook(path,read_only=True,data_only=True); names=book.sheetnames; book.close()
        for name in names:
            if grid_case_id(name) in done: continue
            feature,audit=grid_summary(path,name); features.append(feature); audits.append(audit); pd.DataFrame(features).sort_values("case_id").to_csv(fp,index=False); pd.DataFrame(audits).sort_values("case_id").to_csv(ap,index=False)
    table=pd.DataFrame(features).sort_values("case_id").drop_duplicates("case_id",keep="last"); audits=pd.DataFrame(audits).sort_values("case_id").drop_duplicates("case_id",keep="last")
    if table.case_id.nunique()!=100: raise ValueError(f"Expected 100 cases, found {table.case_id.nunique()}")
    for parameter in UNCERTAINTIES.values():
        duplicate_columns=[column for column in (f"{parameter}_x",f"{parameter}_y") if column in table]
        if duplicate_columns:
            table[parameter]=table[duplicate_columns].bfill(axis=1).iloc[:,0]
            table=table.drop(columns=duplicate_columns)
    missing_parameters=set(UNCERTAINTIES.values())-set(table.columns)
    if missing_parameters:
        table=table.merge(pd.read_csv(output_path(processed,"case_uncertainties.csv")),on="case_id",how="left",validate="one_to_one")
    if table.isna().any().any(): raise ValueError(f"Features contain missing values: {table.columns[table.isna().any()].tolist()}")
    table.to_csv(fp,index=False); audits.to_csv(ap,index=False); return table,audits

def run(input_dir: Path,processed: Path,with_grid_features=False,resume_grid_features=False) -> dict[str,pd.DataFrame]:
    """Run the full Person 1 pipeline and return each produced audit/table in memory."""
    core=run_core(input_dir,processed); curves=run_curves(input_dir,processed); grid,audit=manifest(input_dir,processed); result={"core_audit":core,"curve_audit":curves,"grid_manifest":grid,"grid_schema_audit":audit}
    if with_grid_features: features,feature_audit=extract_features(input_dir,processed,resume_grid_features); result.update(grid_features=features,grid_feature_audit=feature_audit)
    return result

def main() -> None:
    """Parse command-line options and execute the requested Person 1 workflow."""
    parser=argparse.ArgumentParser(description="Run complete Person 1 data engineering."); parser.add_argument("--input-dir",type=Path,default=ROOT/"TestCases"); parser.add_argument("--processed-dir",type=Path,default=ROOT/"data"/"processed"); parser.add_argument("--with-grid-features",action="store_true"); parser.add_argument("--resume-grid-features",action="store_true"); args=parser.parse_args(); output=run(args.input_dir,args.processed_dir,args.with_grid_features,args.resume_grid_features); print("\n".join(f"{k}: {len(v)} rows" for k,v in output.items()))
if __name__=="__main__": main()
