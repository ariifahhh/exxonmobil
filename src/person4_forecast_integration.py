"""Person 4: produce, validate, and export the provisional 2008-2028 submission workbook.

The saved Person 2 proxy is deliberately restricted to the historical source
window.  This module therefore validates that artifact and the Person 3 handover,
then extrapolates the ten Person 3 curve-ranked cases using the best historical
back-test configuration.  It never asks the proxy to predict beyond 2008.
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.person2_ml_proxy import FLUIDS, PARAMETERS, predict_production
from src.person4_baseline_forecast import DEFAULT_FORECAST_CONFIG, forecast_case


FORECAST_CUTOFF = pd.Timestamp("2008-01-01")
REQUIRED_SHEETS = [f"Case {number}" for number in range(1, 11)]
FLUID_TO_COLUMN = {"gas": 2, "oil": 3, "water": 4}


def load_template_dates(template_path: Path) -> pd.DatetimeIndex:
    """Read and validate the shared 80-date schedule from the untouched template."""
    dates = pd.DatetimeIndex(pd.to_datetime(
        pd.read_excel(template_path, sheet_name="Case 1", usecols="A", engine="openpyxl").iloc[:, 0], errors="raise"
    ))
    if len(dates) != 80 or dates.has_duplicates or not dates.is_monotonic_increasing:
        raise ValueError("Template must contain 80 unique, ascending forecast dates.")
    if dates.min() <= FORECAST_CUTOFF:
        raise ValueError("Template forecast dates must begin after the 2008-01-01 historical cutoff.")
    return dates


def selected_cases(processed_dir: Path, top_n: int = 10) -> pd.DataFrame:
    """Load exactly ten Person 3 curve-ranked cases and their Person 1 parameters."""
    ranking = pd.read_csv(processed_dir / "person3_curve_history_ranking.csv")
    parameters = pd.read_csv(processed_dir / "person1_case_uncertainties.csv")
    selected = ranking.head(top_n).merge(parameters[["case_id", *PARAMETERS]], on="case_id", validate="one_to_one")
    if len(selected) != top_n or selected["case_id"].nunique() != top_n:
        raise ValueError(f"Expected {top_n} unique Person 3 ranked cases.")
    selected.insert(0, "forecast_number", np.arange(1, top_n + 1))
    return selected


def load_forecast_config(processed_dir: Path) -> dict[str, float]:
    """Select the lowest-error configuration from the fixed historical back-test."""
    summary_path = processed_dir / "person4_forecast_backtest_summary.csv"
    if not summary_path.exists():
        raise FileNotFoundError("Run src/person4_forecast_backtest.py before final integration.")
    best = pd.read_csv(summary_path).sort_values(["weighted_nrmse", "config"]).iloc[0]
    config = DEFAULT_FORECAST_CONFIG.copy()
    config.update({key: best[key] for key in config if key in best.index})
    config["name"] = str(best["config"])
    return config


def integration_contract_audit(processed_dir: Path, model_path: Path) -> pd.DataFrame:
    """Load Person 2/3 artifacts and record that their machine-readable contract works."""
    optimizer_path = processed_dir / "person3_optimized_parameters.csv"
    config_path = processed_dir / "person3_optimizer_config.json"
    metadata_path = model_path.with_name("person2_production_proxy_metadata.json")
    if not model_path.exists() or not metadata_path.exists() or not optimizer_path.exists() or not config_path.exists():
        raise FileNotFoundError("Person 2 model or Person 3 optimizer handover artifact is missing.")

    optimizer = pd.read_csv(optimizer_path).iloc[0]
    optimizer_config = json.loads(config_path.read_text(encoding="utf-8"))
    model_metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    parameter_values = {name: float(optimizer[name]) for name in PARAMETERS}
    supported_end = pd.Timestamp(model_metadata["supported_date_end"])
    proxy_rows = predict_production(parameter_values, [supported_end], model_path)
    bounds = optimizer_config["bounds_from_person1_case_range"]
    within_bounds = all(float(bounds[name][0]) <= parameter_values[name] <= float(bounds[name][1]) for name in PARAMETERS)
    rows = [
        ("person2_saved_proxy_loads_without_retraining", len(proxy_rows), len(FLUIDS), "pass" if len(proxy_rows) == len(FLUIDS) else "fail"),
        ("person2_proxy_output_schema", set(proxy_rows.columns) == {"date", "fluid", "predicted_rate"}, True, "pass" if set(proxy_rows.columns) == {"date", "fluid", "predicted_rate"} else "fail"),
        ("person2_proxy_rates_non_negative", bool((proxy_rows["predicted_rate"] >= 0).all()), True, "pass" if (proxy_rows["predicted_rate"] >= 0).all() else "fail"),
        ("person3_parameter_order_matches_proxy", optimizer_config["parameter_order"] == PARAMETERS, True, "pass" if optimizer_config["parameter_order"] == PARAMETERS else "fail"),
        ("person3_optimum_within_person1_bounds", within_bounds, True, "pass" if within_bounds else "fail"),
        ("post_2008_method", "backtested_rate_extrapolation", "backtested_rate_extrapolation", "pass"),
    ]
    return pd.DataFrame(rows, columns=["check", "observed", "expected", "status"])


def build_forecast_table(
    processed_dir: Path, dates: pd.DatetimeIndex, config: dict[str, float]
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Extrapolate each selected P3 case while retaining its actual 2008 cumulative anchor."""
    curves = pd.read_csv(processed_dir / "person1_case_production_quarterly.csv", parse_dates=["date"])
    rows: list[pd.DataFrame] = []
    parameter_rows: list[dict[str, object]] = []
    for candidate in selected_cases(processed_dir).itertuples(index=False):
        history = curves.loc[(curves["case_id"] == candidate.case_id) & (curves["date"] <= FORECAST_CUTOFF)].copy()
        if history.groupby("fluid").size().to_dict() != {fluid: 41 for fluid in FLUIDS}:
            raise ValueError(f"Case {candidate.case_id} does not contain the complete 1998-2008 production history.")
        forecast, fitted = forecast_case(history, dates, config)
        anchors = history.loc[history["date"] == FORECAST_CUTOFF, ["fluid", "date", "rate", "cumulative"]].rename(
            columns={"date": "cutoff_date", "rate": "cutoff_rate", "cumulative": "cutoff_cumulative"}
        )
        forecast = forecast.merge(anchors, on="fluid", validate="many_to_one")
        forecast.insert(0, "forecast_number", candidate.forecast_number)
        forecast.insert(1, "case_id", candidate.case_id)
        forecast["history_match_rank"] = candidate.rank
        forecast["history_match_score"] = candidate.history_match_score
        rows.append(forecast)
        for fitted_row in fitted:
            parameter_rows.append({
                "forecast_number": candidate.forecast_number,
                "case_id": candidate.case_id,
                "history_match_rank": candidate.rank,
                "history_match_score": candidate.history_match_score,
                "forecast_config": config["name"],
                **{name: getattr(candidate, name) for name in PARAMETERS},
                **fitted_row,
            })
    return pd.concat(rows, ignore_index=True), pd.DataFrame(parameter_rows)


def validate_final_forecast(forecast: pd.DataFrame, template_dates: pd.DatetimeIndex) -> pd.DataFrame:
    """Audit schedule, physical plausibility, boundary anchor, and trapezoidal integration."""
    checks: list[dict[str, object]] = []
    for (number, case_id, fluid), group in forecast.groupby(["forecast_number", "case_id", "fluid"], sort=True):
        group = group.sort_values("date").reset_index(drop=True)
        prior_dates = pd.concat([pd.Series([group.loc[0, "cutoff_date"]]), group["date"].iloc[:-1]], ignore_index=True)
        prior_rates = pd.concat([pd.Series([group.loc[0, "cutoff_rate"]]), group["rate"].iloc[:-1]], ignore_index=True).to_numpy(float)
        prior_cumulative = np.r_[float(group.loc[0, "cutoff_cumulative"]), group["cumulative"].iloc[:-1].to_numpy(float)]
        expected_delta = ((prior_rates + group["rate"].to_numpy(float)) / 2) * (group["date"].to_numpy() - prior_dates.to_numpy()) / np.timedelta64(1, "D")
        integration_error = float(np.max(np.abs(group["cumulative"].to_numpy(float) - (prior_cumulative + expected_delta))))
        expected_dates_match = pd.DatetimeIndex(group["date"]).equals(template_dates)
        checks.append({
            "forecast_number": number,
            "case_id": case_id,
            "fluid": fluid,
            "expected_dates": len(template_dates),
            "actual_dates": len(group),
            "template_dates_match": expected_dates_match,
            "duplicate_dates": int(group["date"].duplicated().sum()),
            "unsorted_dates": int(not group["date"].is_monotonic_increasing),
            "missing_values": int(group[["rate", "cumulative"]].isna().sum().sum()),
            "negative_rates": int((group["rate"] < 0).sum()),
            "decreasing_cumulative": int((group["cumulative"].diff() < -1e-8).sum()),
            "boundary_below_cutoff": int(float(group.loc[0, "cumulative"]) < float(group.loc[0, "cutoff_cumulative"])),
            "max_integration_abs_error": integration_error,
        })
    audit = pd.DataFrame(checks)
    audit["status"] = np.where(
        (audit["actual_dates"] == audit["expected_dates"])
        & audit["template_dates_match"]
        & (audit[["duplicate_dates", "unsorted_dates", "missing_values", "negative_rates", "decreasing_cumulative", "boundary_below_cutoff"]].sum(axis=1) == 0)
        & (audit["max_integration_abs_error"] < 1e-5),
        "pass",
        "fail",
    )
    return audit


def export_template(template_path: Path, output_path: Path, forecast: pd.DataFrame, dates: pd.DatetimeIndex) -> None:
    """Copy Template 09, fill only B:D values, and preserve its sheets and layout."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)
    workbook = openpyxl.load_workbook(output_path)
    if workbook.sheetnames != REQUIRED_SHEETS:
        raise ValueError("Template sheet names/layout differ from the required Case 1 through Case 10 sheets.")
    for number, sheet_name in enumerate(REQUIRED_SHEETS, start=1):
        sheet = workbook[sheet_name]
        selected = forecast.loc[forecast["forecast_number"] == number]
        cumulative = selected.pivot(index="date", columns="fluid", values="cumulative").reindex(dates)
        for row_offset, (_, row) in enumerate(cumulative.iterrows(), start=2):
            for fluid, column in FLUID_TO_COLUMN.items():
                sheet.cell(row=row_offset, column=column).value = float(row[fluid])
    workbook.save(output_path)


def validate_exported_workbook(output_path: Path, forecast: pd.DataFrame, dates: pd.DatetimeIndex) -> pd.DataFrame:
    """Reopen the export and verify all ten sheets retain dates and exact forecast values."""
    workbook = openpyxl.load_workbook(output_path, data_only=False)
    rows = []
    for number, sheet_name in enumerate(REQUIRED_SHEETS, start=1):
        sheet = workbook[sheet_name]
        sheet_dates = pd.DatetimeIndex(pd.to_datetime([sheet.cell(row=index, column=1).value for index in range(2, 82)]))
        expected = forecast.loc[forecast["forecast_number"] == number].pivot(index="date", columns="fluid", values="cumulative").reindex(dates)
        values_match = True
        for offset, (_, row) in enumerate(expected.iterrows(), start=2):
            values_match &= all(np.isclose(float(sheet.cell(row=offset, column=column).value), float(row[fluid])) for fluid, column in FLUID_TO_COLUMN.items())
        rows.append({"sheet": sheet_name, "date_count": len(sheet_dates), "dates_match": sheet_dates.equals(dates), "values_match": values_match, "status": "pass" if sheet_dates.equals(dates) and values_match else "fail"})
    return pd.DataFrame(rows)


def run(
    input_dir: Path = ROOT / "TestCases",
    processed_dir: Path = ROOT / "data" / "processed",
    model_path: Path = ROOT / "models" / "person2_production_proxy.joblib",
    output_dir: Path = ROOT / "outputs",
) -> dict[str, pd.DataFrame | Path]:
    """Execute the complete provisional Person 4 integration and workbook export."""
    template_path = input_dir / "09 Template Deliverable.xlsx"
    dates = load_template_dates(template_path)
    contracts = integration_contract_audit(processed_dir, model_path)
    config = load_forecast_config(processed_dir)
    forecast, parameters = build_forecast_table(processed_dir, dates, config)
    audit = validate_final_forecast(forecast, dates)
    if (contracts["status"] != "pass").any() or (audit["status"] != "pass").any():
        raise ValueError("Integration validation failed; workbook was not exported.")
    delivery_path = output_dir / "person4_delivery" / "Template Deliverable - provisional forecast.xlsx"
    export_template(template_path, delivery_path, forecast, dates)
    workbook_audit = validate_exported_workbook(delivery_path, forecast, dates)
    if (workbook_audit["status"] != "pass").any():
        raise ValueError("Workbook export validation failed.")
    forecast.to_csv(processed_dir / "person4_final_forecast_long.csv", index=False, date_format="%Y-%m-%d")
    parameters.to_csv(processed_dir / "person4_final_forecast_parameters.csv", index=False)
    audit.to_csv(processed_dir / "person4_final_forecast_audit.csv", index=False)
    contracts.to_csv(processed_dir / "person4_integration_contract_audit.csv", index=False)
    workbook_audit.to_csv(processed_dir / "person4_template_export_audit.csv", index=False)
    return {"forecast": forecast, "parameters": parameters, "audit": audit, "contracts": contracts, "workbook_audit": workbook_audit, "delivery_path": delivery_path}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run the provisional Person 4 forecast and Template 09 export.")
    parser.add_argument("--input-dir", type=Path, default=ROOT / "TestCases")
    parser.add_argument("--processed-dir", type=Path, default=ROOT / "data" / "processed")
    parser.add_argument("--model-path", type=Path, default=ROOT / "models" / "person2_production_proxy.joblib")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "outputs")
    return parser.parse_args()


if __name__ == "__main__":
    arguments = parse_args()
    result = run(arguments.input_dir, arguments.processed_dir, arguments.model_path, arguments.output_dir)
    print(result["contracts"].to_string(index=False))
    print(result["audit"].to_string(index=False))
    print(f"\nCreated and verified: {result['delivery_path']}")
